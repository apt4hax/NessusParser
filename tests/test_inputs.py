"""CLI regression tests; run: python -m unittest discover -s tests -v"""
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / 'nessus2excel.py'


def scan(index):
    return f'''<NessusClientData_v2><Report name="test"><ReportHost name="192.0.2.{index}">
    <HostProperties><tag name="hostname">host-{index}</tag></HostProperties>
    <ReportItem severity="3" pluginName="High finding"><synopsis>Example</synopsis></ReportItem>
    <ReportItem severity="0" pluginName="Info finding"/>
    </ReportHost></Report></NessusClientData_v2>'''


class InputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.scans = self.root / 'customer scans'
        self.scans.mkdir()
        self.paths = []
        for i in range(1, 17):
            path = self.scans / f'{i:02}.nessus'
            path.write_text(scan(i))
            self.paths.append(path)
        (self.scans / 'ignored.txt').write_text('not XML')
        nested = self.scans / 'nested'
        nested.mkdir()
        (nested / 'nested.nessus').write_text(scan(99))
        self.output = self.root / 'output.xlsx'

    def run_cli(self, *args, ok=True):
        result = subprocess.run([sys.executable, str(SCRIPT), *map(str, args), '-o',
                                 str(self.output)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0 if ok else 2, result.stderr)
        return result

    def rows(self):
        wb = load_workbook(self.output)
        try:
            self.assertEqual(wb.sheetnames, ['Nessus Findings'])
            ws = wb.active
            self.assertEqual(ws.max_column, 10)
            self.assertEqual(ws.freeze_panes, 'A2')
            return list(ws.values)
        finally:
            wb.close()

    def test_directory_matches_16_explicit_files(self):
        explicit = self.run_cli(*self.paths)
        expected = self.rows()
        directory = self.run_cli('-dir', self.scans)
        self.assertEqual(self.rows(), expected)
        self.assertEqual(len(expected), 17)
        self.assertEqual(explicit.stdout, directory.stdout)

    def test_overlap_glob_and_repeated_directories(self):
        self.run_cli(self.paths[0], str(self.scans / '*.nessus'), '-dir', self.scans,
                     '--dir', self.scans)
        self.assertEqual(len(self.rows()), 17)

    def test_single_file_xml_and_severity(self):
        xml = self.root / 'scan.xml'
        xml.write_text(scan(1))
        self.run_cli(xml)
        self.assertEqual(len(self.rows()), 2)
        self.run_cli(xml, '--min-severity', 0)
        self.assertEqual(len(self.rows()), 3)

    def test_uppercase_extension(self):
        self.paths[0].rename(self.scans / '01.NESSUS')
        self.run_cli('--dir', self.scans)
        self.assertEqual(len(self.rows()), 17)

    def test_invalid_inputs_preserve_existing_output(self):
        empty = self.root / 'empty'
        empty.mkdir()
        bad = self.root / 'bad.nessus'
        bad.write_text('<broken')
        for args in [(), ('-dir', empty), ('-dir', self.root / 'missing'),
                     ('-dir', bad), (self.root / 'missing.nessus',),
                     (self.paths[0], bad)]:
            with self.subTest(args=args):
                self.output.write_bytes(b'existing report')
                self.run_cli(*args, ok=False)
                self.assertEqual(self.output.read_bytes(), b'existing report')


if __name__ == '__main__':
    unittest.main()
