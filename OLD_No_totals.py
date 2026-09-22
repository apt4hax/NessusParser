#!/usr/bin/env python3
"""
Parse Tenable Nessus XML (.nessus/.xml) files into an Excel workbook.

Output columns:
Host IP, Host Name, Operating System, Severity, Synopsis, Description,
Solution, Plugin Output, See Also, CVE Numbers

Usage:
  python nessus_to_excel.py scan1.nessus scan2.xml -o nessus_findings.xlsx
  python nessus_to_excel.py scans/*.nessus --min-severity 1 -o nessus_findings.xlsx
"""

from __future__ import annotations

import argparse
import glob
import ipaddress
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Iterable, List, Dict, Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit("Missing dependency: openpyxl. Install it with: pip install openpyxl") from exc


COLUMNS = [
    "Host IP",
    "Host Name",
    "Operating System",
    "Severity",
    "Synopsis",
    "Description",
    "Solution",
    "Plugin Output",
    "See Also",
    "CVE Numbers",
]

SEVERITY_MAP = {
    "0": "Info",
    "1": "Low",
    "2": "Medium",
    "3": "High",
    "4": "Critical",
}

SEVERITY_ORDER = {"Info": 0, "Low": 1, "Medium": 2, "High": 3, "Critical": 4}


def text_or_blank(parent: ET.Element, tag_name: str) -> str:
    node = parent.find(tag_name)
    return (node.text or "").strip() if node is not None else ""


def multiline_child_values(parent: ET.Element, tag_name: str) -> str:
    values = []
    for node in parent.findall(tag_name):
        value = (node.text or "").strip()
        if value:
            values.append(value)
    return "\n".join(values)


def comma_child_values(parent: ET.Element, tag_name: str) -> str:
    values = []
    for node in parent.findall(tag_name):
        value = (node.text or "").strip()
        if value:
            values.append(value)
    return ", ".join(sorted(set(values)))


def host_property(host: ET.Element, names: Iterable[str]) -> str:
    props = host.find("HostProperties")
    if props is None:
        return ""

    wanted = {name.lower() for name in names}
    for tag in props.findall("tag"):
        tag_name = (tag.attrib.get("name") or "").lower()
        if tag_name in wanted:
            value = (tag.text or "").strip()
            if value:
                return value
    return ""


def looks_like_ip(value: str) -> bool:
    try:
        ipaddress.ip_address(value)
        return True
    except ValueError:
        return False


def severity_label(raw: str) -> str:
    raw = (raw or "").strip()
    return SEVERITY_MAP.get(raw, raw.title() if raw else "Info")


def expand_input_paths(patterns: Iterable[str]) -> List[Path]:
    paths: List[Path] = []
    for pattern in patterns:
        matches = glob.glob(pattern)
        if matches:
            paths.extend(Path(m) for m in matches)
        else:
            paths.append(Path(pattern))

    unique = []
    seen = set()
    for path in paths:
        resolved = path.resolve()
        if resolved not in seen:
            seen.add(resolved)
            unique.append(path)
    return unique


def parse_nessus_file(path: Path, min_severity: int = 1) -> List[Dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path}")

    tree = ET.parse(path)
    root = tree.getroot()
    rows: List[Dict[str, Any]] = []

    for host in root.findall(".//ReportHost"):
        report_host_name = (host.attrib.get("name") or "").strip()

        host_ip = host_property(host, ["host-ip", "ipv4-address", "host-ipv4"])
        if not host_ip and looks_like_ip(report_host_name):
            host_ip = report_host_name

        host_name = host_property(
            host,
            ["host-fqdn", "hostname", "netbios-name", "dns-name", "host-rdns", "host-hostname"],
        )
        if not host_name and report_host_name and not looks_like_ip(report_host_name):
            host_name = report_host_name

        operating_system = host_property(
            host,
            ["operating-system", "os", "system-type", "cpe-0"],
        )

        for item in host.findall("ReportItem"):
            sev_num = int(item.attrib.get("severity", "0") or 0)

            # Exclude informational findings by default.
            if sev_num < min_severity:
                continue

            synopsis = text_or_blank(item, "synopsis") or item.attrib.get("pluginName", "")
            description = text_or_blank(item, "description")
            solution = text_or_blank(item, "solution")
            plugin_output = text_or_blank(item, "plugin_output")
            see_also = multiline_child_values(item, "see_also")
            cve_numbers = comma_child_values(item, "cve")

            rows.append(
                {
                    "Host IP": host_ip,
                    "Host Name": host_name,
                    "Operating System": operating_system,
                    "Severity": severity_label(str(sev_num)),
                    "Synopsis": synopsis.strip(),
                    "Description": description.strip(),
                    "Solution": solution.strip(),
                    "Plugin Output": plugin_output.strip(),
                    "See Also": see_also,
                    "CVE Numbers": cve_numbers,
                }
            )

    return rows


def autosize_columns(ws) -> None:
    min_widths = {
        "Host IP": 14,
        "Host Name": 18,
        "Operating System": 22,
        "Severity": 10,
        "Synopsis": 20,
        "Description": 30,
        "Solution": 30,
        "Plugin Output": 30,
        "See Also": 24,
        "CVE Numbers": 18,
    }

    max_widths = {
        "Host IP": 18,
        "Host Name": 28,
        "Operating System": 40,
        "Severity": 12,
        "Synopsis": 45,
        "Description": 80,
        "Solution": 80,
        "Plugin Output": 90,
        "See Also": 50,
        "CVE Numbers": 40,
    }

    for col_idx, column_name in enumerate(COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        max_len = 0

        for cell in ws[letter]:
            value = str(cell.value or "")
            max_len = max(max_len, min(len(value), max_widths[column_name]))

        ws.column_dimensions[letter].width = max(
            min_widths[column_name],
            min(max_len + 2, max_widths[column_name]),
        )


def write_excel(rows: List[Dict[str, Any]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Nessus Findings"

    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(col, "") for col in COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    severity_fills = {
        "Critical": "C00000",
        "High": "FF0000",
        "Medium": "FFC000",
        "Low": "92D050",
        "Info": "D9EAF7",
    }

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(COLUMNS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border

        severity_cell = row[COLUMNS.index("Severity")]
        severity = severity_cell.value
        if severity in severity_fills:
            severity_cell.fill = PatternFill("solid", fgColor=severity_fills[severity])
            severity_cell.font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    autosize_columns(ws)
    ws.sheet_view.showGridLines = False
    wb.save(output_path)


def main(argv: List[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Convert Nessus XML output to Excel.")
    parser.add_argument("inputs", nargs="+", help="Nessus .nessus/.xml file paths or glob patterns")
    parser.add_argument("-o", "--output", default="nessus_findings.xlsx", help="Output Excel file path")
    parser.add_argument(
        "--min-severity",
        type=int,
        choices=[0, 1, 2, 3, 4],
        default=1,
        help="Minimum numeric severity to include: 0=Info, 1=Low, 2=Medium, 3=High, 4=Critical",
    )
    args = parser.parse_args(argv)

    all_rows: List[Dict[str, Any]] = []
    for path in expand_input_paths(args.inputs):
        all_rows.extend(parse_nessus_file(path, args.min_severity))

    all_rows.sort(
        key=lambda r: (
            -SEVERITY_ORDER.get(r.get("Severity", "Info"), 0),
            r.get("Host IP", ""),
            r.get("Host Name", ""),
            r.get("Synopsis", ""),
        )
    )

    write_excel(all_rows, Path(args.output))
    print(f"Wrote {len(all_rows)} findings to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())